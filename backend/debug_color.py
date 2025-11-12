# --- Cópialo y guárdalo como debug_color.py ---
from __future__ import annotations
import re
from typing import Dict, Any, Optional, Tuple
from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX

# --- INICIO: Funciones copiadas de tu sso_transformer.py ---
# (Las necesitamos para resolver los colores de TEMA)

def _normalize_argb(s: str) -> Optional[str]:
    if not s: return None
    s = s.strip().upper()
    if s in ("00000000", "FF000000", "NONE", "AUTO"): return None
    if s.startswith("#"): s = s[1:]
    if re.fullmatch(r"[0-9A-F]{6}", s): return "FF" + s
    if re.fullmatch(r"[0-9A-F]{8}", s): return s
    return s

def _rgb_to_tuple(rgb6: str) -> Tuple[int, int, int]:
    rgb6 = rgb6.lstrip("#").upper()
    return int(rgb6[0:2], 16), int(rgb6[2:4], 16), int(rgb6[4:6], 16)

def _tuple_to_rgb6(t: Tuple[int, int, int]) -> str:
    r, g, b = (max(0, min(255, v)) for v in t)
    return f"{r:02X}{g:02X}{b:02X}"

def _apply_tint(rgb6: str, tint: Optional[float]) -> str:
    if tint is None or tint == 0: return rgb6
    r, g, b = _rgb_to_tuple(rgb6)
    if tint < 0:
        factor = 1.0 + tint
        r, g, b = int(r * factor), int(g * factor), int(b * factor)
    else:
        r = int(r + (255 - r) * tint)
        g = int(g + (255 - g) * tint)
        b = int(b + (255 - b) * tint)
    return _tuple_to_rgb6((r, g, b))

def _build_theme_map(wb) -> Dict[int, str]:
    defaults = {
        0: "FFFFFF", 1: "000000", 2: "EEECE1", 3: "1F497D",
        4: "4F81BD", 5: "C0504D", 6: "9BBB59", 7: "8064A2",
        8: "4BACC6", 9: "F79646", 10: "0000FF", 11: "800080",
    }
    try:
        xml = wb.theme._element
        clrScheme = None
        for child in xml.iterchildren():
            if child.tag.endswith("themeElements"):
                for el in child.iterchildren():
                    if el.tag.endswith("clrScheme"):
                        clrScheme = el
                        break
        if clrScheme is None: return defaults
        def _extract_rgb(node) -> Optional[str]:
            for ch in node.iterchildren():
                tag = ch.tag.split("}")[-1]
                if tag == "srgbClr": return ch.get("val")
                if tag == "sysClr": return ch.get("lastClr")
            return None
        keys = ["lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3",
                "accent4", "accent5", "accent6", "hlink", "folHlink"]
        out: Dict[int, str] = {}
        idx = 0
        for node in clrScheme.iterchildren():
            tag = node.tag.split("}")[-1]
            if tag in keys:
                rgb = _extract_rgb(node)
                if rgb: out[idx] = rgb.upper()
                else: out[idx] = defaults[idx]
                idx += 1
        for i in range(12): out.setdefault(i, defaults.get(i, "000000"))
        return out
    except Exception:
        return defaults

def _get_cell_color_hex(cell, theme_map: Dict[int, str]) -> Optional[str]:
    fill = getattr(cell, "fill", None)
    if fill is None: return None
    ptype = getattr(fill, "patternType", None) or getattr(fill, "fill_type", None)
    if ptype not in ("solid", "darkGray", "lightGray", "gray125", "gray0625"):
        return None
    fg = getattr(fill, "fgColor", None) or getattr(fill, "start_color", None)
    if fg is None: return None
    ctype = getattr(fg, "type", None)
    if ctype == "rgb":
        return _normalize_argb(getattr(fg, "rgb", None))
    if ctype == "indexed":
        idx = getattr(fg, "index", None)
        try:
            if idx is not None and 0 <= int(idx) < len(COLOR_INDEX):
                return _normalize_argb(COLOR_INDEX[int(idx)].upper().lstrip("#"))
        except Exception: pass
    if ctype == "theme":
        theme = getattr(fg, "theme", None)
        tint = getattr(fg, "tint", None)
        try: base = theme_map.get(int(theme))
        except Exception: base = None
        if base:
            rgb6 = _apply_tint(base, float(tint) if tint is not None else None)
            return _normalize_argb(rgb6)
        return f"THEME:{theme}:{tint}"
    return None

# --- FIN: Funciones copiadas de tu sso_transformer.py ---


# --- SCRIPT PRINCIPAL DE PRUEBA ---

# 1. DEFINE TUS DATOS
FILE_PATH = '/Users/diegoleivareyes/Documents/UAI/POD_C006/POD N°193 - Obras Tempranas PRET - Movimiento de Tierra Masivo Rev0 - 24-07-25.xlsx'
SHEET_NAME = '3.SSO'

# Celdas que tienen los NÚMEROS de los días
# (Según tus logs y capturas)
DAY_CELLS_TO_CHECK = {
    "Día 7 (Verde)": "C11",
    "Día 25 (Amarillo)": "I15",
    "Día 28 (Blanco)": "I17",
}

print(f"--- Abriendo archivo: {FILE_PATH} ---")
try:
    wb = load_workbook(FILE_PATH, data_only=True)
    ws = wb[SHEET_NAME]
    theme_map = _build_theme_map(wb)
    print("--- ¡Listos! Probando tu teoría ---")

    for test_name, cell_ref in DAY_CELLS_TO_CHECK.items():
        print(f"\n--- {test_name} ---")
        
        # 1. Revisa la celda con el NÚMERO
        cell_num = ws[cell_ref]
        color_num = _get_cell_color_hex(cell_num, theme_map)
        print(f"  Celda del NÚMERO ({cell_ref}):\t\tValor='{cell_num.value}', Color={color_num}")

        # 2. Revisa la celda de la IZQUIERDA
        try:
            cell_left = ws.cell(row=cell_num.row, column=cell_num.column - 1)
            color_left = _get_cell_color_hex(cell_left, theme_map)
            print(f"  Celda de la IZQUIERDA ({cell_left.coordinate}):\tValor='{cell_left.value}', Color={color_left}")
        except Exception:
            print(f"  No se pudo leer la celda de la izquierda de {cell_ref}")

except FileNotFoundError:
    print(f"ERROR: ¡Archivo no encontrado! Verifica la ruta:")
    print(FILE_PATH)
except Exception as e:
    print(f"Ocurrió un error inesperado: {e}")