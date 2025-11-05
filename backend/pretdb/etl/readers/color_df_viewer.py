# backend/pretdb/etl/readers/color_df_viewer.py
# -*- coding: utf-8 -*-
"""
Visualiza una hoja de Excel como DataFrame coloreado y lo exporta a HTML.

Uso:
  python backend/pretdb/etl/readers/color_df_viewer.py \
      --path "/ruta/archivo.xlsx" \
      --sheet "3.SSO" \
      --out "/ruta/salida/sso_debug.html" \
      --max-rows 120 --max-cols 80

Requisitos:
  pip install pandas openpyxl
"""

from __future__ import annotations
import argparse
import os
from typing import Optional, Tuple, List

import pandas as pd
from openpyxl import load_workbook


# --------------------------- utils de color ---------------------------

def _argb_to_hex(argb: Optional[str]) -> Optional[str]:
    """Convierte 'FF00B050' o '00B050' a '#00B050'."""
    if not argb:
        return None
    s = argb.strip()
    if len(s) == 8:  # ARGB → usa los últimos 6
        return f"#{s[2:].upper()}"
    if len(s) == 6:  # RGB
        return f"#{s.upper()}"
    return None


def _extract_hex_from_rgb_obj(obj) -> Optional[str]:
    """
    Algunas versiones de openpyxl devuelven un objeto RGB en lugar de un str.
    Intentamos extraer un string desde sus posibles atributos.
    """
    if obj is None:
        return None

    # Caso 1: ya es string
    if isinstance(obj, str):
        return _argb_to_hex(obj)

    # Caso 2: tiene atributos comunes con el string dentro
    for attr in ("rgb", "srgb", "value"):
        val = getattr(obj, attr, None)
        if isinstance(val, str):
            return _argb_to_hex(val)

    # Último recurso: str(obj) podría venir como 'FF00B050' en algunas impls
    s = str(obj)
    s = s.replace("'", "").replace('"', "").strip()
    hx = _argb_to_hex(s)
    return hx


def _color_to_hex(color) -> Optional[str]:
    """
    Recibe un objeto Color de openpyxl y devuelve '#RRGGBB' si es posible.
    Maneja variantes de openpyxl (string u objeto RGB intermedio).
    """
    if color is None:
        return None

    # Intento directo
    rgb_attr = getattr(color, "rgb", None)
    hx = _extract_hex_from_rgb_obj(rgb_attr)
    if hx:
        return hx

    # A veces hay 'index', 'theme', etc. No se resuelven aquí.
    # Retorna None si no hay un rgb explícito.
    return None


def _cell_bg_hex(cell) -> Optional[str]:
    """
    Extrae el color de fondo de una celda openpyxl como '#RRGGBB' si el relleno es sólido.
    Si no hay color o el relleno no es sólido, devuelve None.
    """
    fill = getattr(cell, "fill", None)
    if not fill or getattr(fill, "patternType", None) != "solid":
        return None

    # Prioridad: fgColor → start_color → end_color (según disponibilidad)
    for attr in ("fgColor", "start_color", "end_color"):
        col = getattr(fill, attr, None)
        if col is None:
            continue
        hx = _color_to_hex(col)
        if hx:
            return hx
    return None


# --------------------------- carga de datos ---------------------------

def load_values_df(xlsx_path: str, sheet_name: str, max_rows: int, max_cols: int) -> pd.DataFrame:
    """Lee valores con pandas (sin estilos)."""
    df = pd.read_excel(
        xlsx_path, sheet_name=sheet_name, header=None, dtype=object, engine="openpyxl"
    )
    if max_rows > 0:
        df = df.iloc[:max_rows, :]
    if max_cols > 0:
        df = df.iloc[:, :max_cols]
    return df


def load_colors_df(xlsx_path: str, sheet_name: str, shape: Tuple[int, int]) -> pd.DataFrame:
    """Lee colores con openpyxl, devolviendo un DataFrame paralelo (mismas dimensiones)."""
    wb = load_workbook(xlsx_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"La hoja '{sheet_name}' no existe en el archivo.")
    ws = wb[sheet_name]

    nrows, ncols = shape
    colors = []
    for r in range(1, nrows + 1):
        row_colors = []
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            row_colors.append(_cell_bg_hex(cell))
        colors.append(row_colors)
    return pd.DataFrame(colors)


# --------------------------- estilizado HTML ---------------------------

def style_with_colors(df_values: pd.DataFrame, df_colors: pd.DataFrame) -> pd.io.formats.style.Styler:
    """Aplica los colores de df_colors al Styler de df_values."""
    def row_style(row: pd.Series) -> List[str]:
        r = row.name
        styles: List[str] = []
        for c in range(len(row)):
            hexc = df_colors.iat[r, c]
            styles.append(f"background-color: {hexc}" if isinstance(hexc, str) else "")
        return styles

    styler = (
        df_values
        .style
        .apply(row_style, axis=1)
        .hide(axis="index")
        .set_table_attributes('border="1" cellpadding="4" cellspacing="0"')
    )
    return styler


def build_legend_html(df_colors: pd.DataFrame) -> str:
    """Crea una pequeña leyenda con los colores únicos encontrados y sus frecuencias."""
    s = df_colors.stack(dropna=False)
    s = s.map(lambda x: x if isinstance(x, str) else "None")
    counts = s.value_counts()
    parts = ['<div style="margin:10px 0;font-family:system-ui,Arial;">',
             '<strong>Colores únicos encontrados (conteo de celdas):</strong><br>']
    for hexc, cnt in counts.items():
        if hexc == "None":
            parts.append(
                '<span style="display:inline-block;width:14px;height:14px;'
                'border:1px solid #ccc;vertical-align:middle;background:white"></span> '
                f'None: {cnt}&nbsp;&nbsp;'
            )
        else:
            parts.append(
                f'<span style="display:inline-block;width:14px;height:14px;'
                f'border:1px solid #000;vertical-align:middle;background:{hexc}"></span> '
                f'{hexc}: {cnt}&nbsp;&nbsp;'
            )
    parts.append("</div>")
    return "".join(parts)


def save_html(styler: pd.io.formats.style.Styler, legend_html: str, out_path: str, title: str) -> None:
    """Guarda el HTML final (leyenda + tabla)."""
    table_html = styler.to_html()
    html = f"""<!doctype html>
<html>
<head>
  <meta charset="utf-8" />
  <title>{title}</title>
</head>
<body style="margin:24px;font-family:system-ui,Arial;">
  <h2 style="margin-top:0">{title}</h2>
  {legend_html}
  {table_html}
</body>
</html>
"""
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html)


# --------------------------- CLI ---------------------------

def main():
    ap = argparse.ArgumentParser(description="Renderiza una hoja Excel a HTML con colores de fondo por celda.")
    ap.add_argument("--path", required=True, help="Ruta al archivo .xlsx")
    ap.add_argument("--sheet", required=True, help="Nombre de la hoja (ej: '3.SSO')")
    ap.add_argument("--out", required=True, help="Ruta del HTML de salida")
    ap.add_argument("--max-rows", type=int, default=0, help="Máximo de filas (0 = todas las detectadas por pandas)")
    ap.add_argument("--max-cols", type=int, default=0, help="Máximo de columnas (0 = todas las detectadas por pandas)")
    args = ap.parse_args()

    df_vals = load_values_df(args.path, args.sheet, args.max_rows, args.max_cols)
    df_cols = load_colors_df(args.path, args.sheet, df_vals.shape)

    styler = style_with_colors(df_vals, df_cols)
    legend = build_legend_html(df_cols)
    title = f"Vista coloreada — {os.path.basename(args.path)} :: {args.sheet} ({df_vals.shape[0]}x{df_vals.shape[1]})"

    save_html(styler, legend, args.out, title)
    print(f"[OK] HTML generado: {args.out}")


if __name__ == "__main__":
    main()