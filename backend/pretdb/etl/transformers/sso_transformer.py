# pretdb/etl/transformers/sso_transformer.py
from __future__ import annotations
from typing import Dict, Any, List, Optional, Tuple
import logging, re
from datetime import datetime, date, timedelta
from calendar import month_name

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX

from pretdb.etl.transformers.registry import register_transformer

from django.db import transaction, connection
from pretdb.models import Pod, Sso, SsoCruzSeguridad, SsoEstadoDia, SsoReflexion

logger = logging.getLogger(__name__)

# ----------------------------- configuración -----------------------------

MESES: Dict[str, int] = {
    "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
    "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "OCTUBRE": 10, "NOVIEMBRE": 11, "DICIEMBRE": 12,
}
MESES_INV: Dict[int, str] = {v: k.title() for k, v in MESES.items()}

# Ajusta según tus colores reales (ya puestos con los códigos que detectaste)
COLOR_A_ESTADO: Dict[str, int] = {
    "FF00B050": 0,  # Verde
    "FFFFFF00": 1,  # Amarillo
    "FFFF0000": 2,  # Rojo
}

# ------------------------------ utilidades ------------------------------

def _normalize_argb(s: str) -> Optional[str]:
    """
    Normaliza un color a ARGB de 8 dígitos (sin '#').
    Devuelve None para vacíos/transparente/negro común.
    """
    if not s:
        return None
    s = s.strip().upper()
    if s in ("00000000", "FF000000", "NONE", "AUTO"):
        return None
    if s.startswith("#"):
        s = s[1:]
    if re.fullmatch(r"[0-9A-F]{6}", s):  # RGB
        return "FF" + s
    if re.fullmatch(r"[0-9A-F]{8}", s):  # ARGB
        return s
    return s

def _rgb_to_tuple(rgb6: str) -> Tuple[int, int, int]:
    rgb6 = rgb6.lstrip("#").upper()
    return int(rgb6[0:2], 16), int(rgb6[2:4], 16), int(rgb6[4:6], 16)

def _tuple_to_rgb6(t: Tuple[int, int, int]) -> str:
    r, g, b = (max(0, min(255, v)) for v in t)
    return f"{r:02X}{g:02X}{b:02X}"

def _apply_tint(rgb6: str, tint: Optional[float]) -> str:
    """
    Aplica el algoritmo de tint de Excel a un RGB6.
    tint > 0 aclara, tint < 0 oscurece.
    """
    if tint is None or tint == 0:
        return rgb6
    r, g, b = _rgb_to_tuple(rgb6)
    if tint < 0:
        factor = 1.0 + tint  # tint in [-1,0)
        r = int(r * factor)
        g = int(g * factor)
        b = int(b * factor)
    else:
        r = int(r + (255 - r) * tint)
        g = int(g + (255 - g) * tint)
        b = int(b + (255 - b) * tint)
    return _tuple_to_rgb6((r, g, b))

def _build_theme_map(wb) -> Dict[int, str]:
    """
    Extrae el esquema de colores del theme del workbook.
    Devuelve {theme_index -> RGB6}. Si falla, entrega defaults de Office.
    """
    defaults = {
        0: "FFFFFF", 1: "000000", 2: "EEECE1", 3: "1F497D",
        4: "4F81BD", 5: "C0504D", 6: "9BBB59", 7: "8064A2",
        8: "4BACC6", 9: "F79646", 10: "0000FF", 11: "800080",
    }
    try:
        th = wb.theme
        if th is None:
            return defaults
        xml = th._element  # CT_Theme

        clrScheme = None
        for child in xml.iterchildren():
            if child.tag.endswith("themeElements"):
                for el in child.iterchildren():
                    if el.tag.endswith("clrScheme"):
                        clrScheme = el
                        break
        if clrScheme is None:
            return defaults

        def _extract_rgb(node) -> Optional[str]:
            for ch in node.iterchildren():
                tag = ch.tag.split("}")[-1]
                if tag == "srgbClr":
                    val = ch.get("val")
                    if val:
                        return val.upper()
                if tag == "sysClr":
                    val = ch.get("lastClr")
                    if val:
                        return val.upper()
            return None

        keys = [
            "lt1", "dk1", "lt2", "dk2",
            "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
            "hlink", "folHlink",
        ]
        out: Dict[int, str] = {}
        idx = 0
        for node in clrScheme.iterchildren():
            tag = node.tag.split("}")[-1]
            if tag in keys:
                rgb = _extract_rgb(node) or defaults[idx]
                out[idx] = rgb
                idx += 1
        for i in range(12):
            out.setdefault(i, defaults.get(i, "000000"))
        return out
    except Exception:
        return defaults

def _get_cell_color_hex(cell, theme_map: Dict[int, str]) -> Optional[str]:
    """
    Devuelve el ARGB (8 dígitos) del relleno si es sólido.
    Soporta fgColor.type: 'rgb', 'indexed', 'theme' (+ tint).
    """
    fill = getattr(cell, "fill", None)
    if fill is None:
        return None

    ptype = getattr(fill, "patternType", None) or getattr(fill, "fill_type", None)
    if ptype not in ("solid", "darkGray", "lightGray", "gray125", "gray0625"):
        return None

    fg = getattr(fill, "fgColor", None) or getattr(fill, "start_color", None)
    if fg is None:
        return None

    ctype = getattr(fg, "type", None)

    # 1) RGB directo
    if ctype == "rgb":
        raw = getattr(fg, "rgb", None)
        if isinstance(raw, str):
            return _normalize_argb(raw)

    # 2) INDEXED -> usa COLOR_INDEX
    if ctype == "indexed":
        idx = getattr(fg, "index", None)
        try:
            if idx is not None:
                idx = int(idx)
                if 0 <= idx < len(COLOR_INDEX):
                    rgb6 = COLOR_INDEX[idx].upper().lstrip("#")
                    return _normalize_argb(rgb6)
        except Exception:
            pass

    # 3) THEME -> convierte con theme_map + tint
    if ctype == "theme":
        theme = getattr(fg, "theme", None)
        tint = getattr(fg, "tint", None)
        try:
            base = theme_map.get(int(theme)) if theme is not None else None
        except Exception:
            base = None
        if base:
            rgb6 = _apply_tint(base, float(tint) if tint is not None else None)
            return _normalize_argb(rgb6)
        return f"THEME:{theme}:{tint}"

    # 4) fallback a end_color
    endc = getattr(fill, "end_color", None)
    if endc is not None:
        ctype2 = getattr(endc, "type", None)
        if ctype2 == "rgb":
            raw = getattr(endc, "rgb", None)
            if isinstance(raw, str):
                return _normalize_argb(raw)
        if ctype2 == "indexed":
            idx = getattr(endc, "index", None)
            try:
                if idx is not None:
                    idx = int(idx)
                    if 0 <= idx < len(COLOR_INDEX):
                        rgb6 = COLOR_INDEX[idx].upper().lstrip("#")
                        return _normalize_argb(rgb6)
            except Exception:
                pass
        if ctype2 == "theme":
            theme = getattr(endc, "theme", None)
            tint = getattr(endc, "tint", None)
            try:
                base = theme_map.get(int(theme)) if theme is not None else None
            except Exception:
                base = None
            if base:
                rgb6 = _apply_tint(base, float(tint) if tint is not None else None)
                return _normalize_argb(rgb6)
            return f"THEME:{theme}:{tint}"

    return None

def _mes_nombre(mes: Optional[int]) -> Optional[str]:
    if mes is None:
        return None
    mes_int = int(mes)
    if mes_int in MESES_INV:
        return MESES_INV[mes_int]
    if 1 <= mes_int <= 12:
        return month_name[mes_int].title()
    return None

def _obtener_mes_y_anio(ws, df: pd.DataFrame) -> Tuple[Optional[int], int]:
    mes_num: Optional[int] = None
    for row in ws.iter_rows(min_row=1, max_row=15, values_only=True):
        for value in row:
            if isinstance(value, str):
                txt = value.strip().upper()
                if txt in MESES:
                    mes_num = MESES[txt]
                    break
        if mes_num is not None:
            break
    if mes_num is None:
        logger.warning("SSO: no se pudo detectar el mes en la hoja; se usará fallback.")

    anio: Optional[int] = None
    for val in df.to_numpy().ravel():
        if isinstance(val, (pd.Timestamp, datetime, date)):
            anio = pd.to_datetime(val).year
            break
    if anio is None:
        anio = date.today().year
    return mes_num, anio

def _extraer_dias_cruz(ws, theme_map: Dict[int, str]) -> List[Dict[str, Any]]:
    """
    Lee celdas con valores 1..31 y toma el color desde la celda a la IZQUIERDA
    (si existe); si no, usa la misma celda.
    """
    dias: List[Dict[str, Any]] = []
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, (int, float)) and 1 <= int(v) <= 31:
                dia = int(v)
                celda_color = cell
                if cell.column > 1:
                    try:
                        celda_color = ws.cell(row=cell.row, column=cell.column - 1)
                    except Exception as e:
                        logger.warning(f"No se pudo leer celda a la izquierda de {cell.coordinate}: {e}")
                        celda_color = cell
                color_hex = _get_cell_color_hex(celda_color, theme_map)
                estado = COLOR_A_ESTADO.get(color_hex)
                dias.append({
                    "dia": dia, "fila": cell.row, "columna": cell.column,
                    "color": color_hex, "estado": estado
                })
    return dias

def _buscar_valor_por_etiqueta(df: pd.DataFrame, etiqueta: str) -> Optional[float]:
    """
    Busca fila que contenga 'etiqueta' y retorna el primer número en esa fila.
    (DataFrame.applymap para la máscara; Series.map para filtrar numéricos).
    """
    mask = df.applymap(lambda x: isinstance(x, str) and etiqueta.lower() in x.lower())
    coords = list(zip(*mask.to_numpy().nonzero()))
    if not coords:
        return None
    fila_idx, _ = coords[0]
    fila = df.loc[fila_idx]
    nums = fila[fila.map(lambda x: isinstance(x, (int, float)))]
    if nums.empty:
        return None
    try:
        return float(nums.iloc[0])
    except Exception:
        return None

def _month_bounds(y: int, m: int) -> Tuple[date, date]:
    start = date(y, m, 1)
    if m == 12:
        end = date(y + 1, 1, 1) - timedelta(days=1)
    else:
        end = date(y, m + 1, 1) - timedelta(days=1)
    return start, end

# ----------------------------- clase principal -----------------------------

@register_transformer("sso")
class SSOTransformer:
    """
    Transformer para SSO (Cruz de Seguridad + Hallazgos/Tarjeta/Policlínico).
    Solo persiste el estado del día de la fecha del POD (con fallback ±3d).
    """
    sheet_key = "sso"

    def run_df(
        self,
        df_raw: pd.DataFrame,
        meta: Dict[str, Any],
        dry_run: bool = False,
        **kwargs: Any,
    ) -> Dict[str, Any]:
        logger.debug("SSOTransformer.run_df: shape=%s", df_raw.shape)
        df = df_raw.copy()

        file_path = (
            kwargs.get("file_path")
            or meta.get("source_path")
            or meta.get("file_path")
        )
        if not file_path:
            raise ValueError(
                "SSOTransformer necesita 'file_path' o 'source_path' en meta/kwargs "
                "para leer colores de la cruz."
            )

        sheet_name = (
            kwargs.get("sheet_name")
            or meta.get("raw_sheet_name")
            or meta.get("sheet_key")
        )

        wb = load_workbook(file_path, data_only=True)
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            logger.warning("SSOTransformer: sheet '%s' no encontrada; usando activa.", sheet_name)
            ws = wb.active

        theme_map = _build_theme_map(wb)
        mes, anio = _obtener_mes_y_anio(ws, df)
        dias = _extraer_dias_cruz(ws, theme_map)

        hallazgos = _buscar_valor_por_etiqueta(df, "Hallazgos")
        tarjeta_verde = _buscar_valor_por_etiqueta(df, "Tarjeta Verde")
        policlinico = _buscar_valor_por_etiqueta(df, "Policlinico")

        summary = {
            "mes": mes,
            "anio": anio,
            "mes_nombre": _mes_nombre(mes),
            "n_dias_detectados": len(dias),
            "n_dias_con_estado": sum(1 for d in dias if d.get("estado") is not None),
            "hallazgos": hallazgos,
            "tarjeta_verde": tarjeta_verde,
            "policlinico": policlinico,
        }

        return {
            "sheet": self.sheet_key,
            "rows": len(dias),
            "summary": summary,
            "dias_detail": dias,
            "dry_run": dry_run,
        }

    @transaction.atomic
    def persist(self, result: dict, run_id: int) -> Tuple[int, Dict, Optional[Dict]]:
        # Delegate to the new loader implementation in loaders.sso_loader.
        try:
            from pretdb.etl.loaders.sso_loader import SsoLoader
        except Exception as e:
            logger.exception("No se pudo importar el loader de SSO: %s", e)
            return 0, {"import": str(e)}, None

        pod = self._get_pod_obj(run_id)
        if not pod:
            return 0, {"pod_id": "etl_run.pod_id es NULL; ejecuta 'portada' antes."}, None

        loader = SsoLoader()
        return loader.persist(pod, result, dry_run=False)

    # --------------------------- helpers internos ---------------------------

    def _upsert_reflexion_only(self, pod: Pod, fecha: date, summ: Dict[str, Any], warnings: Dict[str, Any]) -> None:
        hall = summ.get("hallazgos")
        tv = summ.get("tarjeta_verde")
        poli = summ.get("policlinico")
        if hall is None and tv is None and poli is None:
            return
        try:
            sso, _ = Sso.objects.get_or_create(pod=pod)
            SsoReflexion.objects.update_or_create(
                sso=sso,
                fecha=fecha,
                defaults={
                    "hallazgos": int(hall) if hall is not None else None,
                    "tarjeta_verde": int(tv) if tv is not None else None,
                    "traslado_policlinico": int(poli) if poli is not None else None,
                }
            )
        except Exception as e:
            warnings["reflexion"] = f"No se pudo upsert SsoReflexion: {e}"

    def _get_pod_obj(self, run_id: int) -> Optional[Pod]:
        with connection.cursor() as cur:
            cur.execute("SELECT pod_id FROM etl_run WHERE id=%s", [run_id])
            row = cur.fetchone()
        if not row or not row[0]:
            return None
        try:
            return Pod.objects.get(id=row[0])
        except Pod.DoesNotExist:
            return None

# ----------------------- Debug para colores reales -----------------------

def debug_colores_dias(path_excel: str, sheet_name: Optional[str] = None) -> None:
    """
    Imprime colores únicos en celdas con días 1..31, usando la lógica de
    mirar a la izquierda del número (si existe).
    """
    wb = load_workbook(path_excel, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    theme_map = _build_theme_map(wb)

    print(f"--- Depurando colores de {path_excel} (Hoja: {ws.title}) ---")
    print("--- Usando lógica: 'color de la celda a la izquierda del número' ---")
    vistos = set()
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, (int, float)) and 1 <= int(v) <= 31:
                celda_color = cell
                if cell.column > 1:
                    try:
                        celda_color = ws.cell(row=cell.row, column=cell.column - 1)
                    except Exception:
                        celda_color = cell
                color_hex = _get_cell_color_hex(celda_color, theme_map)
                if color_hex not in vistos:
                    vistos.add(color_hex)
                    print(f"Celda {cell.coordinate} (día={int(v)}): Color en {celda_color.coordinate} -> {color_hex}")

    if not vistos:
        print("SSO: no se encontraron colores en días 1..31.")
    print("--- Fin de la depuración ---")

# ----------------------- Wrapper p/orquestador -----------------------

def load_sso(df: pd.DataFrame, run_id: int, sheet: str, **kwargs):
    """
    Orquestador: transforma y persiste. kwargs: file_path=..., sheet_name=...
    Retorna (rows_in, rows_upserted, warnings, errors)
    """
    tr = SSOTransformer()
    meta = {
        "raw_sheet_name": sheet,
        **({"file_path": kwargs.get("file_path")} if kwargs.get("file_path") else {})
    }
    result = tr.run_df(df, meta=meta, dry_run=False, **kwargs)
    rows_in = result.get("rows", 0)
    try:
        rows_upserted, warnings, errors = tr.persist(result, run_id=run_id)
        return rows_in, rows_upserted, warnings, errors
    except Exception as e:
        logger.exception("Error en persistencia de SSO: %s", e)
        return rows_in, 0, {}, {"persist": str(e)}
